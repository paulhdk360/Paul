"""Génère le Bon de livraison (BL) 2027 d'Enedril.
Usage : python3 generer_bl.py <dossier_sortie> [exemple]"""
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.filters import FilterColumn, CustomFilters, CustomFilter
from openpyxl.utils import get_column_letter as L

OUTDIR = sys.argv[1]; EXEMPLE = len(sys.argv) > 2
LOGO = '/home/user/Paul/public/logo-enedril-white.png'
NAVY, BLUE, CYAN, RED = '08203F', '0A4FA0', '00AEEF', 'E2231A'
INK, MUTED, LINE, INPUT = '0B1B2E', '51647A', 'D5DEE8', 'EAF6FD'
F = 'Calibri'
font = lambda sz=10, b=False, c=INK, i=False: Font(name=F, size=sz, bold=b, color=c, italic=i)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = lambda c=LINE: Side(style='thin', color=c)

NL = 24                       # lignes du BL (comme l'import)
I1 = 6                        # première ligne de l'import
T1 = 15                       # première ligne du tableau du BL
USED = 4

wb = Workbook(); ws = wb.active; ws.title = '1 Bon de livraison'
wi = wb.create_sheet('2 Import devis'); wl = wb.create_sheet('Listes')
for w, c in ((ws, BLUE), (wi, CYAN), (wl, '9FB6CE')): w.sheet_properties.tabColor = c; w.sheet_view.showGridLines = False
wl['A1'] = 'Type de colisage'; wl['A1'].font = font(9, True, 'FFFFFF'); wl['A1'].fill = fill(NAVY); wl.column_dimensions['A'].width = 24
for i, n in enumerate(['Palette', 'Fishing basket', 'Caisse', 'Vrac'], 2): wl[f'A{i}'] = n
def dv(rng, formula, strict=False):
    d = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=strict); ws.add_data_validation(d); d.add(rng)

# ------------------------------------------------------------------ 2 import
for col, w in zip('ABCDEFGHI', (3, 60, 7, 10, 10, 10, 12, 10, 10)): wi.column_dimensions[col].width = w
for r in (1, 2, 3):
    for c in range(1, 10): wi.cell(r, c).fill = fill(NAVY if r < 3 else CYAN)
wi.row_dimensions[1].height = 30; wi.row_dimensions[2].height = 20; wi.row_dimensions[3].height = 5
wi['B1'] = 'IMPORT DU DEVIS'; wi['B1'].font = font(18, True, 'FFFFFF'); wi['B1'].alignment = Alignment(vertical='center')
wi['B2'] = 'Dans le classeur Devis, onglet « 2 Matériel » : copier les colonnes C à J (lignes 6 à 29) et les coller ici en B6. Seuls la description, la quantité et la valeur de remplacement passent sur le BL.'
wi['B2'].font = font(10, c='9FD8F5'); wi['B2'].alignment = Alignment(vertical='center')
for i, h in enumerate(['', 'Description', 'Qté', 'Stand-by', 'Opér.', 'UC', 'Valeur remplac.\n(LIH / DBR)', 'Inspection', 'Restocking'], 1):
    if i == 1: continue
    c = wi.cell(5, i); c.value = h; c.font = font(9, True, 'FFFFFF'); c.fill = fill(BLUE if i in (2, 3, 7) else '7A8CA3'); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wi.row_dimensions[5].height = 32
for r in range(I1, I1 + NL):
    for col in 'BCDEFGHI':
        c = wi[f'{col}{r}']; c.fill = fill(INPUT); c.font = font(10, False, BLUE); c.border = Border(bottom=Side(style='thin', color=CYAN))
        c.number_format = '#,##0.00' if col in 'DEFGHI' else 'General'; c.alignment = Alignment(horizontal='left' if col == 'B' else 'center', vertical='center', indent=1 if col == 'B' else 0)
wi['B31'] = 'Colonnes grises : tarifs, non repris sur le BL (le BL peut être envoyé au transporteur et au client sans les prix de location).'
wi['B31'].font = font(9, i=True, c=MUTED)
if EXEMPLE:
    for i, row in enumerate([('BHA #1', None, None, None, None, None, None, None), ('Stab 5-7/8 OD, 2-1/4" ID - 3-1/2" IF Pin x Box', 1, 85, 120, 40, 4200, 150, 600),
                             ('XO 3-1/2" IF Pin x 2-7/8" EU Box', 2, 20, 30, 10, 900, 60, 150), ('XO 3-1/2" IF Box x 2-7/8" EU Pin', 1, 20, 30, 10, 900, 60, 150)]):
        for col, v in zip('BCDEFGHI', row):
            if v is not None: wi[f'{col}{I1+i}'] = v

# ------------------------------------------------------------------ 1 BL
for col, w in zip('ABCDEFG', (17, 20, 52, 7, 9, 12, 14)): ws.column_dimensions[col].width = w
ws.column_dimensions['I'].hidden = True
for r in (1, 2, 3):
    for c in range(1, 8): ws.cell(r, c).fill = fill(NAVY if r < 3 else CYAN)
ws.row_dimensions[1].height = 34; ws.row_dimensions[2].height = 22; ws.row_dimensions[3].height = 5
img = XLImage(LOGO); img.width, img.height = 150, 34; ws.add_image(img, 'A1')
ws.merge_cells('C1:G1'); ws['C1'] = 'BON DE LIVRAISON · DELIVERY TICKET'; ws['C1'].font = font(15, True, 'FFFFFF'); ws['C1'].alignment = Alignment(horizontal='right', vertical='center')
ws.merge_cells('C2:G2'); ws['C2'] = '="N° "&A6'; ws['C2'].font = font(11, c='9FD8F5'); ws['C2'].alignment = Alignment(horizontal='right', vertical='center')
ws.merge_cells('A4:G4'); ws['A4'] = 'ENEDRIL SAS · 3 Avenue des Lacs · 64140 Lons (France) · +33 (0)5 59 04 50 45 · SIRET 79464246200020 · TVA FR11794642462'
ws['A4'].font = font(8, c=MUTED); ws['A4'].alignment = Alignment(horizontal='center', vertical='center'); ws.row_dimensions[4].height = 18
def tin(r, c1, c2, text, fmt=None, tall=False):
    if c2 > c1: ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2); ws.merge_cells(start_row=r+1, start_column=c1, end_row=r+1, end_column=c2)
    a = ws.cell(r, c1); a.value = text.upper(); a.font = font(7.5, True, MUTED); a.alignment = Alignment(vertical='bottom', indent=1)
    for k in range(c1, c2 + 1): ws.cell(r+1, k).fill = fill(INPUT); ws.cell(r+1, k).border = Border(bottom=Side(style='thin', color=CYAN))
    b = ws.cell(r+1, c1); b.font = font(10.5, True, BLUE); b.alignment = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
    if fmt: b.number_format = fmt
    ws.row_dimensions[r].height = 15; ws.row_dimensions[r+1].height = 34 if tall else 21
tin(5, 1, 2, 'N° du BL · Delivery ticket n°'); tin(5, 3, 3, 'Date', 'dd/mm/yyyy'); tin(5, 4, 7, 'N° de dossier (devis) · Job number')
tin(7, 1, 2, 'Client · Customer'); tin(7, 3, 3, 'Puits / chantier · Well'); tin(7, 4, 7, 'PO client · Purchase order')
tin(9, 1, 2, 'Transporteur · Transport company'); tin(9, 3, 3, 'PO transport · Transport PO n°'); tin(9, 4, 7, 'Délai · Expected time of arrival (ETA)')
tin(11, 1, 3, 'Adresse de chargement · Loading address', None, True); tin(11, 4, 7, 'Adresse de livraison · Delivery address', None, True)
ws['A12'] = 'Enedril · 3 Avenue des Lacs · 64140 Lons (France)'
ws.row_dimensions[13].height = 8
HD = ['Réf\nRef', 'N° de série\nSerial number', 'Description', 'Qté\nQty', 'Origine\nOrigin', 'Poids net\nkg / unité', 'Valeur de remplacement\n€ / unité']
for i, h in enumerate(HD, 1):
    c = ws.cell(14, i); c.value = h; c.font = font(8.5, True, NAVY); c.fill = fill('DCE8F5'); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[14].height = 36; ws['I14'] = 'actif'; ws['I14'].font = font(7, c='B8C4D2')
last = T1 + NL - 1
for i in range(NL):
    r = T1 + i; m = I1 + i; ws.row_dimensions[r].height = 26
    ws[f'I{r}'] = f"=IF('2 Import devis'!$B{m}=\"\",0,1)"; ws[f'I{r}'].font = font(7, c='B8C4D2')
    ws[f'C{r}'] = f"=IF('2 Import devis'!$B{m}=\"\",\"\",'2 Import devis'!$B{m})"
    ws[f'D{r}'] = f"=IF(OR('2 Import devis'!$B{m}=\"\",N('2 Import devis'!$C{m})=0),\"\",'2 Import devis'!$C{m})"
    ws[f'G{r}'] = f"=IF(OR(D{r}=\"\",'2 Import devis'!$G{m}=\"\"),\"\",'2 Import devis'!$G{m})"
    ws[f'C{r}'].font = font(10.5, c=INK); ws[f'C{r}'].alignment = Alignment(vertical='center', wrap_text=True, indent=1)
    for col in 'ABEF':
        c = ws[f'{col}{r}']; c.fill = fill(INPUT); c.font = font(10.5, False, BLUE); c.alignment = Alignment(horizontal='left' if col in 'AB' else 'center', vertical='center', indent=1 if col in 'AB' else 0, wrap_text=True)
        c.border = Border(bottom=Side(style='thin', color='B7DDF2'))
    ws[f'E{r}'] = 'FR'; ws[f'F{r}'].number_format = '#,##0.0'
    for col in 'DG': ws[f'{col}{r}'].font = font(10.5, c=INK); ws[f'{col}{r}'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'G{r}'].number_format = '#,##0.00'
    for col in 'CDG': ws[f'{col}{r}'].border = Border(bottom=thin())
ws.conditional_formatting.add(f'E{T1}:E{last}', FormulaRule(formula=[f'AND($I{T1}=1,$D{T1}="")'], fill=fill(NAVY), font=Font(name=F, color=NAVY), stopIfTrue=True))
ws.conditional_formatting.add(f'A{T1}:G{last}', FormulaRule(formula=[f'AND($I{T1}=1,$D{T1}="")'], fill=fill(NAVY), font=Font(name=F, bold=True, color='FFFFFF'), stopIfTrue=True))
# lignes vides masquées par un filtre ; Ctrl+Alt+L (Réappliquer) pour mettre à jour
ws.auto_filter.ref = f'A14:I{last}'
for cid in range(9):
    ws.auto_filter.filterColumn.append(FilterColumn(colId=cid, hiddenButton=True, showButton=False, customFilters=CustomFilters(customFilter=[CustomFilter(operator='greaterThan', val='0')]) if cid == 8 else None))
for i in range(USED, NL): ws.row_dimensions[T1 + i].hidden = True
ws.merge_cells('A13:G13'); ws['A13'] = f'=IF(SUBTOTAL(103,C{T1}:C{last})<SUM(I{T1}:I{last}),"⚠  Appuyer sur Ctrl + Alt + L pour afficher toutes les lignes du BL.","")'
ws['A13'].font = font(9, True, RED); ws['A13'].alignment = Alignment(indent=1, vertical='center'); ws.row_dimensions[13].height = 14
# totaux
rt = last + 1
ws.merge_cells(start_row=rt, start_column=1, end_row=rt, end_column=3)
ws[f'A{rt}'] = 'TOTAL'; ws[f'D{rt}'] = f'=SUM(D{T1}:D{last})'
ws[f'F{rt}'] = f'=SUMPRODUCT(N(+D{T1}:D{last}),N(+F{T1}:F{last}))' if False else f'=SUMPRODUCT(IF(ISNUMBER(D{T1}:D{last}),D{T1}:D{last},0)*IF(ISNUMBER(F{T1}:F{last}),F{T1}:F{last},0))'
ws[f'G{rt}'] = f'=SUMPRODUCT(IF(ISNUMBER(D{T1}:D{last}),D{T1}:D{last},0)*IF(ISNUMBER(G{T1}:G{last}),G{T1}:G{last},0))'
ws[f'F{rt}'].number_format = '#,##0.0 "kg"'; ws[f'G{rt}'].number_format = '#,##0.00 "€"'
for c in range(1, 8): ws.cell(rt, c).fill = fill(NAVY); ws.cell(rt, c).font = font(10.5, True, 'FFFFFF'); ws.cell(rt, c).alignment = Alignment(horizontal='left' if c < 4 else 'center', vertical='center', indent=1 if c == 1 else 0)
ws.row_dimensions[rt].height = 24
ws.row_dimensions[rt + 1].height = 18; ws.row_dimensions[rt + 2].height = 18
# colisage
P0 = rt + 3
ws.merge_cells(start_row=P0, start_column=1, end_row=P0, end_column=7); ws.cell(P0, 1).value = 'COLISAGE · PACKING (ensemble des articles)'
for c in range(1, 8): ws.cell(P0, c).fill = fill(NAVY)
ws.cell(P0, 1).font = font(10, True, 'FFFFFF'); ws.cell(P0, 1).alignment = Alignment(horizontal='center', vertical='center'); ws.row_dimensions[P0].height = 22
for i, h in enumerate(['Colis · Package', 'Type de colisage · Type', 'Dimensions L × l × h (cm)', '', '', 'Poids brut\nkg', 'Volume\nm³'], 1):
    c = ws.cell(P0 + 1, i); c.value = h; c.font = font(8.5, True, NAVY); c.fill = fill('DCE8F5'); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[P0 + 1].height = 30
ws.merge_cells(start_row=P0 + 1, start_column=3, end_row=P0 + 1, end_column=5)
# dimensions : trois cases L, l, h dans C (texte) ? -> colonnes dédiées D, E (étroites) : on place L/l/h dans C (saisie libre "120 × 80 × 60")
NP = 3
for i in range(NP):
    r = P0 + 2 + i; ws.row_dimensions[r].height = 26
    ws[f'A{r}'] = f'Colis {i+1}'; ws[f'A{r}'].font = font(10.5, True, NAVY); ws[f'A{r}'].alignment = Alignment(horizontal='center', vertical='center')
    for col in 'BCFG':
        c = ws[f'{col}{r}']; c.fill = fill(INPUT); c.font = font(10.5, False, BLUE); c.alignment = Alignment(horizontal='center', vertical='center'); c.border = Border(bottom=Side(style='thin', color='B7DDF2'))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws[f'F{r}'].number_format = '#,##0.0'
    ws[f'G{r}'].number_format = '0.000'
dv(f'B{P0+2}:B{P0+1+NP}', '=Listes!$A$2:$A$6')
ws.cell(P0 + 2 + NP, 1).value = '="TOTAL  ·  "&COUNTA(B' + str(P0 + 2) + ':B' + str(P0 + 1 + NP) + ')&" colis"'
ws.merge_cells(start_row=P0 + 2 + NP, start_column=1, end_row=P0 + 2 + NP, end_column=5)
rcol = P0 + 2 + NP
ws[f'F{rcol}'] = f'=SUM(F{P0+2}:F{rcol-1})'; ws[f'F{rcol}'].number_format = '#,##0.0 "kg"'; ws[f'G{rcol}'] = f'=SUM(G{P0+2}:G{rcol-1})'; ws[f'G{rcol}'].number_format = '0.000 "m³"'
for c in range(1, 8): ws.cell(rcol, c).fill = fill(NAVY); ws.cell(rcol, c).font = font(10, True, 'FFFFFF'); ws.cell(rcol, c).alignment = Alignment(horizontal='left' if c == 1 else 'center', vertical='center', indent=1 if c == 1 else 0)
ws.row_dimensions[rcol].height = 22
# signatures
rs = rcol + 2
ws.merge_cells(start_row=rs, start_column=1, end_row=rs, end_column=7)
ws[f'A{rs}'] = 'Marchandises reçues en bon état, à retourner à Enedril à la fin de la location · Goods received in good condition, to be returned to Enedril at the end of the rental'
ws[f'A{rs}'].font = font(8.5, True, MUTED); ws[f'A{rs}'].alignment = Alignment(indent=1, vertical='center', wrap_text=True); ws.row_dimensions[rs].height = 26
for c1, c2, t in ((1, 2, 'EXPÉDITEUR · ENEDRIL\nNom, date, signature'), (3, 3, 'TRANSPORTEUR\nNom, date, signature'), (4, 7, 'RÉCEPTION CLIENT · RECEIVED BY\nNom, date, signature, cachet')):
    ws.merge_cells(start_row=rs + 1, start_column=c1, end_row=rs + 4, end_column=c2)
    ws.cell(rs + 1, c1).value = t; ws.cell(rs + 1, c1).font = font(8, True, MUTED); ws.cell(rs + 1, c1).alignment = Alignment(vertical='top', indent=1, wrap_text=True)
    for rr in range(rs + 1, rs + 5):
        for cc in range(c1, c2 + 1): ws.cell(rr, cc).border = Border(left=thin() if cc == c1 else None, right=thin() if cc == c2 else None, top=thin() if rr == rs + 1 else None, bottom=thin() if rr == rs + 4 else None)
for rr in range(rs + 1, rs + 5): ws.row_dimensions[rr].height = 17
LASTROW = rs + 4

if EXEMPLE:
    ws['A6'] = 'BL-26-0123'; ws['C6'] = datetime.date(2026, 11, 9); ws['D6'] = '26-82-FR'
    ws['A8'] = 'Geopetrol'; ws['C8'] = 'Puits Lacq-12 · CLU108'; ws['D8'] = 'PO-4500123'
    ws['A10'] = 'Transports Durand'; ws['C10'] = 'TPO-778'; ws['D10'] = '48 h après chargement'
    ws['D12'] = 'Geopetrol · Base de Lacq · 64170 Lacq'
    ws[f'A{T1+1}'] = 'PL-0412'; ws[f'B{T1+1}'] = 'SN-4471'; ws[f'F{T1+1}'] = 85
    ws[f'A{T1+2}'] = 'PL-0377'; ws[f'B{T1+2}'] = 'SN-1120 / SN-1121'; ws[f'F{T1+2}'] = 12.5
    ws[f'A{T1+3}'] = 'PL-0378'; ws[f'B{T1+3}'] = 'SN-1130'; ws[f'F{T1+3}'] = 12.5
    ws[f'B{P0+2}'] = 'Fishing basket'; ws[f'C{P0+2}'] = '120 × 80 × 60'; ws[f'F{P0+2}'] = 145; ws[f'G{P0+2}'] = 0.576

ws.print_area = f'A1:G{LASTROW}'; ws.page_setup.orientation = 'portrait'; ws.page_setup.paperSize = 9
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = .4; ws.page_margins.top = .4; ws.page_margins.bottom = .5
ws.oddFooter.center.text = 'Enedril · Bon de livraison &P / &N'; ws.oddFooter.center.size = 8
for w in (wi, wl):
    w.page_setup.orientation = 'landscape'; w.page_setup.paperSize = 9; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0; w.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
out = os.path.join(OUTDIR, 'BL_Enedril_2027' + ('_exemple' if EXEMPLE else '') + '.xlsx'); wb.save(out); print('ok', out)
