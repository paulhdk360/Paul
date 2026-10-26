# Gazette Enedril 2026 : python3 -I generer_gazette.py  (depuis ce dossier)
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
NAVY=C(8,32,63);BLUE=C(10,79,160);CYAN=C(28,178,254);LIGHT=C(234,244,252);WHITE=C(255,255,255);INK=C(11,27,46);YEL=C(255,243,176);GREY=C(90,107,127)
F='Ubuntu Mono';W,H=595,842
P=Presentation();P.slide_width=Pt(W);P.slide_height=Pt(H)
def rect(s,x,y,w,h,fill,line=None,dash=False):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));r.shadow.inherit=False
    if fill is None:r.fill.background()
    else:r.fill.solid();r.fill.fore_color.rgb=fill
    if line:r.line.color.rgb=line;r.line.width=Pt(1);r.line.dash_style=MSO_LINE.DASH if dash else None
    else:r.line.fill.background()
    return r
def txt(s,x,y,w,h,text,size=10,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,line=None):
    tb=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=tb.text_frame;tf.word_wrap=True;tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,t in enumerate(text if isinstance(text,list) else [text]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
        if line:p.line_spacing=Pt(line)
        p.space_after=Pt(size*0.5)
        r=p.add_run();r.text=t;r.font.size=Pt(size);r.font.bold=bold;r.font.name=F;r.font.color.rgb=color
    return tb
def slide(bg=WHITE):
    s=P.slides.add_slide(P.slide_layouts[6]);rect(s,0,0,W,H,bg);return s
def title(s,t,y=34):
    txt(s,40,y,W-80,34,t,24,BLUE,True);rect(s,40,y+36,W-80,4,CYAN)
def foot(s,n,dark=False):
    txt(s,40,H-34,300,12,'Enedril · Rétrospective 2026',8,CYAN if dark else GREY)
    txt(s,W-80,H-34,40,12,str(n),8,CYAN if dark else GREY,align=PP_ALIGN.RIGHT)
def todo(s,x,y,w,h,label):
    rect(s,x,y,w,h,YEL,C(200,160,0),True);txt(s,x+6,y,w-12,h,'À COMPLÉTER : '+label,8.5,C(110,80,0),True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
def bullets(s,x,y,w,items,size=10.5,gap=8):
    for t in items:
        n=max(1,-(-len(t)*size*0.6//w))
        rect(s,x,y+3,6,6,CYAN);txt(s,x+14,y,w-14,n*size*1.35,t,size,INK,line=size*1.3);y+=n*size*1.3+gap
    return y
def stat(s,x,y,w,h,val,label,ph=True):
    rect(s,x,y,w,h,BLUE);txt(s,x,y+8,w,30,val,22,WHITE,True,PP_ALIGN.CENTER);txt(s,x+6,y+h-30,w-12,26,label,8.5,C(190,225,250),align=PP_ALIGN.CENTER)
def photo(s,x,y,w,h,label):
    rect(s,x,y,w,h,C(238,242,247),C(150,170,190),True);txt(s,x+6,y,w-12,h,'PHOTO : '+label,8.5,GREY,False,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
logo='../public/logo-enedril-white.png'
# 1 couverture
s=slide(NAVY);rect(s,300,0,295,H,C(10,60,115));rect(s,0,560,W,282,BLUE)
txt(s,40,150,500,60,'RÉTRO',60,WHITE,True);txt(s,40,215,520,60,'SPECTIVE',60,WHITE,True);txt(s,40,290,300,40,'2026',34,CYAN,True)
s.shapes.add_picture('../outil_SPX_bleu.png',Pt(40),Pt(400),width=Pt(515));s.shapes.add_picture(logo,Pt(40),Pt(740),width=Pt(200))
txt(s,40,690,400,14,'La gazette de fin d\'année · Édition 2026',10,C(190,225,250))
# 2 sommaire
s=slide();title(s,'SOMMAIRE')
for i,(t,p) in enumerate([('Mot de la direction',3),('Actualités internes',4),('Annonces et changements internes',5),('Actualités commerciales',6),('Marketing et communication',7),('Engagements solidaires',8),('Galerie photos',9),('Les jeux',10)]):
    y=130+i*60;rect(s,40,y,515,44,LIGHT);rect(s,40,y,6,44,CYAN);txt(s,60,y,400,44,t,14,NAVY,True,anchor=MSO_ANCHOR.MIDDLE);txt(s,480,y,60,44,str(p),16,BLUE,True,PP_ALIGN.RIGHT,MSO_ANCHOR.MIDDLE)
foot(s,2)
# 3 mot de la direction
s=slide();title(s,'MOT DE LA DIRECTION')
txt(s,40,100,515,640,['Chers collaborateurs,',
'Alors que nous clôturons 2026, nous souhaitons revenir sur les temps forts de cette année et partager avec vous les perspectives qui s\'ouvrent pour la suivante.',
'2026 a été une année de changements concrets pour l\'entreprise : arrivée de nouveaux collaborateurs, essai de la semaine à quatre jours, nouveaux avantages (tickets restaurant et chèques vacances digitaux) et poursuite des travaux sur le site, avec un nouveau bâtiment de stockage, une nouvelle Break-out unit et un parking remanié. Notre système de chauffage a également évolué : une pompe à chaleur alimentée par deux puits de géothermie, une belle cohérence avec notre métier.',
'Côté visibilité, nous avons été présents à GeoTHERM, GEOFLUID, ADIPEC et aux Geodays, et accueilli des étudiants d\'IFP Training. De nouveaux outils sont également arrivés pour compléter notre parc.',
'Nous avons aussi su nous retrouver : teambuilding à la Section Paloise, barbecue en mai, repas après la matinée de sensibilisation aux addictions, sans oublier Octobre rose, Movember et la collecte de jouets pour le Père Noël fait de la récup\'.',
'Nous donnons rendez-vous à tous à GeoTHERM, à Offenburg, les 25 et 26 février 2027, où nous exposerons sur le stand 274.',
'Nous vous remercions pour votre implication et votre sérieux tout au long de l\'année. Votre engagement quotidien est essentiel pour faire progresser Enedril.',
'Merci à tous et à très bientôt,','Fabien & Alexis'],12.5,line=17.5)
foot(s,3)
# 4 actualités internes
s=slide();title(s,'ACTUALITÉS INTERNES')
rect(s,40,100,515,150,LIGHT);rect(s,40,100,6,150,CYAN);txt(s,60,110,480,20,'Bienvenue aux nouveaux',15,BLUE,True)
txt(s,60,138,480,40,'Plusieurs collaborateurs ont rejoint Enedril cette année :',10.5)
txt(s,60,162,480,30,'Thibaut · Mickael · Xavier · Julien Duval · Segen',14,NAVY,True)
todo(s,60,198,480,38,'équipes concernées, départs éventuels à saluer, photo des nouveaux')
txt(s,40,275,515,20,'Événements d\'équipe',15,BLUE,True);rect(s,40,300,515,3,CYAN)
bullets(s,40,315,515,['Teambuilding à la Section Paloise.','Barbecue en mai.','Repas après la matinée de sensibilisation aux addictions.','Repas de fin d\'année (date à confirmer).'])
photo(s,40,470,250,180,'teambuilding Section Paloise');photo(s,305,470,250,180,'barbecue de mai');photo(s,40,665,250,120,'repas / matinée addictions');photo(s,305,665,250,120,'repas de fin d\'année')
foot(s,4)
# 5 annonces
s=slide();title(s,'ANNONCES ET CHANGEMENTS INTERNES')
bullets(s,40,110,515,['Semaine à 4 jours : mise en place d\'un essai.','Tickets restaurant et chèques vacances digitaux : nouveaux avantages mis en place pour les collaborateurs.','Changement du système de chaleur : pompe à chaleur alimentée par les deux puits de géothermie.','Nouveau bâtiment de stockage, nouveau parking et nouvelle Break-out unit.','Arrivée de nouveaux outils.'],12,12)
todo(s,40,380,515,40,'dates, détails des modalités de la semaine à 4 jours, liste des nouveaux outils')
photo(s,40,445,250,170,'nouveau bâtiment de stockage');photo(s,305,445,250,170,'Break-out unit');photo(s,40,630,250,150,'pompe à chaleur / puits');photo(s,305,630,250,150,'nouveaux outils')
foot(s,5)
# 6 commercial
s=slide();title(s,'ACTUALITÉS COMMERCIALES')
txt(s,40,100,515,100,['Cette année, Enedril a accompagné de nombreux projets en géothermie et en forage, en France comme à l\'étranger.','Pour 2027, les perspectives s\'appuient sur le salon GeoTHERM (stand 274) et sur nos échanges en Europe et au Moyen-Orient.'],11,line=15)
for i,(v,l) in enumerate([('…','devis envoyés'),('…','devis pour nouveaux clients'),('…','taux de transformation'),('…','location / vente'),('…','jours de temps de réponse en moyenne')]):
    stat(s,40+(i%3)*175,230+(i//3)*100,165,88,v,l)
for i,(v,l) in enumerate([('…','puits forés'),('…','mètres forés')]):stat(s,215+(i)*175,430,165,88,v,l)
txt(s,40,540,515,20,'Pays d\'intervention',14,BLUE,True);todo(s,40,566,515,40,'liste des pays')
todo(s,40,630,515,60,'statistiques commerciales et de chantier 2026 (les cases « … » ci-dessus) : à fournir par Paul / l\'équipe')
foot(s,6)
# 7 marketing
s=slide();title(s,'MARKETING ET COMMUNICATION')
txt(s,40,100,515,60,'Plusieurs actions ont été menées pour renforcer la visibilité d\'Enedril et valoriser nos activités auprès des partenaires, clients et futurs collaborateurs.',11,line=15)
txt(s,40,170,515,20,'Salons et événements',15,BLUE,True);rect(s,40,194,515,3,CYAN)
bullets(s,40,208,515,['GeoTHERM.','GEOFLUID, Piacenza (Italie), le 9 octobre : Luigi Liberatore représentait Enedril.','ADIPEC, Abu Dhabi (Émirats arabes unis), du 2 au 5 novembre.','Geodays.','Accueil d\'étudiants d\'IFP Training sur nos installations.'],11,8)
photo(s,40,400,165,120,'GeoTHERM');photo(s,215,400,165,120,'GEOFLUID');photo(s,390,400,165,120,'ADIPEC')
txt(s,40,550,515,20,'Côté LinkedIn',15,BLUE,True);rect(s,40,574,515,3,CYAN)
for i,l in enumerate(['posts','abonnés gagnés','abonnés au total','réactions vs 2025']):stat(s,40+i*130,590,122,80,'…',l)
todo(s,40,690,515,40,'statistiques LinkedIn 2026, autres actions marketing (site, presse, vidéo…)')
foot(s,7)
# 8 solidaire
s=slide();title(s,'ENGAGEMENTS SOLIDAIRES')
bullets(s,40,110,515,['Octobre rose : mobilisation de l\'équipe.','Movember : sensibilisation à la santé masculine.','Collecte de jouets : collecte interne pour l\'initiative « Le Père Noël fait de la récup\' ».'],12,12)
todo(s,40,260,515,40,'montants collectés, nombre de jouets, associations bénéficiaires')
photo(s,40,330,250,200,'Octobre rose');photo(s,305,330,250,200,'Movember');photo(s,40,545,515,220,'collecte de jouets')
foot(s,8)
# 9 galerie
s=slide();title(s,'GALERIE PHOTOS')
for i in range(8):photo(s,40+(i%2)*260,100+(i//2)*165,255,155,'à insérer (chantier, atelier, équipe)')
foot(s,9)
# 10 jeux
s=slide();title(s,'LES JEUX')
txt(s,40,100,515,18,'Mots codés (A=1, B=2…)',13,BLUE,True)
def code(w):return '-'.join(str(ord(c)-64) for c in w)
for i,w in enumerate(['FORAGE','GEOTHERMIE','TREPAN','FISHING','MWD','ENEDRIL']):txt(s,40,128+i*22,515,18,code(w)+' : ....................',11)
txt(s,40,275,515,18,'Anagrammes',13,BLUE,True)
ana=[('TUBAGE','GEBUTA'),('SONDAGE','DOSNAGE'),('POMPE','MPEPO'),('PUITS','TIPUS'),('ENERGIE','GENEIRE'),('GAZETTE','TAGEZET')]
for i,(a,b) in enumerate(ana):txt(s,40,303+i*22,515,18,b+' : ....................',11)
txt(s,40,450,515,18,'Devinette',13,BLUE,True)
txt(s,40,476,515,120,['Je descends sous terre, je tourne sans cesse, je mords la roche et je ne m\'arrête jamais avant d\'avoir atteint le fond. Qui suis-je ?','Quelle est la ville du siège d\'Enedril ? Indice : elle est à côté de Pau.'],11,line=15)
rect(s,40,640,515,100,LIGHT);txt(s,52,648,490,16,'Solutions',11,BLUE,True)
txt(s,52,668,490,60,['Mots codés : '+', '.join(['FORAGE','GEOTHERMIE','TREPAN','FISHING','MWD','ENEDRIL']),'Anagrammes : '+', '.join(a for a,b in ana),'Devinette : le trépan · Lons'],9.5,line=12.5)
foot(s,10)
# 11 dos
s=slide(BLUE);rect(s,330,0,265,H,NAVY);s.shapes.add_picture(logo,Pt(60),Pt(370),width=Pt(300))
txt(s,60,440,420,16,'Fishing and Drilling Solutions',11,C(190,225,250));txt(s,60,760,450,14,'www.enedril.com · 3 av des lacs, 64140 Lons',10,WHITE)
P.save('Gazette_Enedril_2026.pptx')
