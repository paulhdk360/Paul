import copy, io
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
IMGONLY={'ImpressionBlock','SafetyJoint','Stabilizer','StringMill','DieCollar','Hook','TaperTap','SuperJunkMill','WashoverShoe','TDog_4-11-16','TDog_5-3-4','TTang_TCatch'}
NAVY_,BLUE_,MID_=None,None,None
FAMILIES=[('CLEANING AND JUNK RECOVERY',NAVY_,[('Scraper','Scraper'),('ImpressionBlock','Impression Block (LIB)'),('CirculatingSub','Circulating Sub'),('RCJB','RCJB'),('JunkBasket','Junk Basket'),('VenturiBasket','Venturi Basket'),('FishingMagnet','Fishing Magnet'),('StringMagnet','String Magnet')]),
 ('CUTTING',BLUE_,[('HEC','Hydraulic External Cutter (HEC)'),('PipeCutter','Pipe Cutter'),('SectionMill','Section Mill')]),
 ('MILLING',MID_,[('Economill','Economill'),('JunkMill','Junk Mill'),('SuperJunkMill','Super Junk Mill'),('ViperPilotMill','Viper Pilot Mill'),('PilotMill','Pilot Mill'),('StringMill','String Mill'),('TaperMill','Taper Mill'),('WashoverShoe','Washover Shoe')]),
 ('HOLE ENLARGEMENT',NAVY_,[('RTU','RTU'),('DTU','DTU'),('HoleOpener','Hole Opener'),('SPX','SPX')]),
 ('JARRING AND IMPACT',BLUE_,[('Intensifier','Intensifier'),('JarZ','Jar Z'),('BumperSub','Bumper Sub'),('DrillingJar','Drilling Jar'),('SafetyJoint','Safety Joint')]),
 ('CATCHING TOOLS',MID_,[('DieCollar','Die Collar'),('Hook','Hook'),('Overshot','Overshot'),('TaperTap','Taper Tap'),('Spear','Spear'),('TTang_TCatch','T-Tang (T catch)'),('TDog_4-11-16','T-Dog 4-11/16'),('TDog_5-3-4','T-Dog 5-3/4')]),
 ('DRILLING AND DIRECTIONAL',NAVY_,[('Stabilizer','Stabilizer (Stab)'),('PDMwithSleeve','PDM with sleeve'),('MWD','Boost MWD'),('Whipstock','Whipstock')])]
order=[k for _,_,l in FAMILIES for k,_ in l]
D='/home/user/Paul/fiches_outils/'
base=Presentation(D+'Fiche_technique_%s.pptx'%order[0])
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def clone_into(prs, src_slide):
    lay=[l for l in prs.slide_layouts if l.name==src_slide.slide_layout.name] or [prs.slide_layouts[0]]
    ns=prs.slides.add_slide(lay[0])
    for sh in list(ns.shapes): sh._element.getparent().remove(sh._element)
    for el in src_slide.shapes._spTree:
        tag=el.tag.split('}')[1]
        if tag in ('nvGrpSpPr','grpSpPr'): continue
        new=copy.deepcopy(el)
        for node in new.iter():
            for att in ('embed','link','id'):
                k='{%s}%s'%(RNS,att)
                if k in node.attrib:
                    rid=node.attrib[k]
                    try: part=src_slide.part.related_part(rid)
                    except KeyError: continue
                    if part.partname.startswith('/ppt/media') or 'image' in part.content_type:
                        _,nr=ns.part.get_or_add_image_part(io.BytesIO(part.blob))
                        node.attrib[k]=nr
        ns.shapes._spTree.append(new)
    # background
    bg=src_slide._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    return ns
from pptx.util import Emu
S='/tmp/claude-0/-home-user-Paul/25d99196-61ed-597d-8b44-fc98194d121f/scratchpad/pdfsl/'
for name in order[1:]:
    if name in IMGONLY:
        lay=base.slide_layouts[6] if len(base.slide_layouts)>6 else base.slide_layouts[0]
        ns=base.slides.add_slide(lay)
        for sh in list(ns.shapes): sh._element.getparent().remove(sh._element)
        ns.shapes.add_picture(S+name+'.png',0,0,width=base.slide_width,height=base.slide_height)
        continue
    src=Presentation(D+'Fiche_technique_%s.pptx'%name)
    clone_into(base,src.slides[0])
exec(open('/tmp/claude-0/-home-user-Paul/25d99196-61ed-597d-8b44-fc98194d121f/scratchpad/catalogue_front.py').read())
base.save('/home/user/Paul/Fiches_outils_Enedril_catalogue.pptx')
print(len(base.slides))
