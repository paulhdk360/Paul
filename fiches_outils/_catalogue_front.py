# s'exécute dans merge_pptx2 : utilise base, FAMILIES, order
import sys
sys.path.insert(0,'/tmp/claude-0/-home-user-Paul/25d99196-61ed-597d-8b44-fc98194d121f/scratchpad')
import build_plaquette as BP
from build_plaquette import rect, text, pic, tile, C, W, H
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
COL={'NAVY':BP.NAVY,'BLUE':BP.BLUE,'MID':BP.MID}
FAMCOL=[BP.NAVY,BP.BLUE,BP.MID,BP.NAVY,BP.BLUE,BP.MID,BP.NAVY]
blank=[l for l in base.slide_layouts if l.name=='Blank'][0]
def newslide():
    ns=base.slides.add_slide(blank)
    for sh in list(ns.shapes): sh._element.getparent().remove(sh._element)
    return ns
page_of={k:i+3 for i,k in enumerate(order)}
# ---- page 1 : presentation
s1=newslide()
rect(s1,0,0,W,300,BP.NAVY)
pic(s1,'logo_white.png',36,34,w=170)
text(s1,36,104,520,[[('TOOL CATALOGUE',dict(size=30,color=BP.WHITE,bold=True))],[('FISHING · MILLING · ENLARGEMENT · JARRING · DRILLING',dict(size=11,color=BP.CYAN,bold=True))]],line=36,after=6)
text(s1,36,190,520,[[('Data sheets of the main downhole tools available from our stock in Lons, near Pau. Each tool has its own page: description, technical characteristics, design features, operating parameters and ordering information.',dict(size=9.5,color=C('DCE9F6')))]],line=13)
for i,(b,sm,f) in enumerate([('2,500','tools in stock',BP.NAVY),('24 / 48h','response time',BP.BLUE),('40','data sheets',BP.MID),('7','tool families',BP.NAVY)]):
    tile(s1,36+i*133,322,125,58,b,sm,f,20)
text(s1,36,404,300,[[('ABOUT ENEDRIL',dict(size=11,color=BP.NAVY,bold=True))]])
rect(s1,36,422,3.5,92,BP.BLUE)
text(s1,50,422,509,[[('Enedril (formerly Mustang Oil Services) is an independent French company founded in 2013 and specialised in directional drilling, fishing and well intervention for geothermal, oil & gas and underground gas storage projects, in France and abroad. Run since 2022 by Fabien and Alexis, the company continues its activities with a new dynamic, and became Enedril in 2025.',dict(size=8.6))]],line=12)
text(s1,36,530,300,[[('HOW TO ORDER A TOOL',dict(size=11,color=BP.NAVY,bold=True))]])
BP.bullets(s1,36,550,523,[[('Find the tool in the summary (next page) or by family. ',False),('Each sheet ends with a "How to order" box.',True)],
   'Tell us the tool type, size or series, top and bottom connections and the casing or hole size.',
   'Dimensions are given in inches and loads in lb unless stated: values are indicative. Other sizes and special versions are available on request.',
   'Fishing response within 24 to 48 hours across metropolitan France.'],size=8.6,line=12,after=4)
rect(s1,36,690,523,92,BP.NAVY)
text(s1,52,690,491,[[('CONTACT',dict(size=8.4,color=BP.CYAN,bold=True))],[('Enedril · 3 avenue des Lacs, 64140 Lons (France)',dict(size=10,color=BP.WHITE,bold=True))],[('info@enedril.fr · +33 (0)5 59 04 50 45 · www.enedril.com',dict(size=10,color=BP.WHITE,bold=True))]],line=14,after=2,anchor=MSO_ANCHOR.MIDDLE,h=92)
BP.footer(s1,1)
# ---- page 2 : summary
s2=newslide()
pic(s2,'logo.png',441,22,w=118)
BP.title(s2,'SUMMARY',y=44)
link_targets=[]
cols=[(36,[0,1,2]),(311,[3,4,5,6])]
for x0,fams in cols:
    y=100
    for fi in fams:
        name,_,items=FAMILIES[fi]
        rect(s2,x0,y,248,22,FAMCOL[fi])
        text(s2,x0+8,y+4,232,[[(name,dict(size=9,color=BP.WHITE,bold=True))]])
        y+=28
        for key,disp in items:
            tb=text(s2,x0+6,y,190,[[(disp,dict(size=8.6,color=BP.INK))]])
            text(s2,x0+200,y,42,[[(str(page_of[key]),dict(size=8.6,color=BP.NAVY,bold=True))]],align=PP_ALIGN.RIGHT)
            rect(s2,x0+6,y+14.5,236,0.5,C('D5E1EE'))
            link_targets.append((tb,key))
            y+=19
        y+=12
BP.footer(s2,2)
# move + link
lst=base.slides._sldIdLst; ids=list(lst)
front=ids[-2:]
for e in front: lst.remove(e)
for k,e in enumerate(front): lst.insert(k,e)
slides=list(base.slides)
for tb,key in link_targets:
    tb.click_action.target_slide=slides[page_of[key]-1]
