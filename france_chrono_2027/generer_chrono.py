"""Génère le France Chrono 2027 d'Enedril (un seul onglet).
Usage : python3 generer_chrono.py <dossier_sortie>"""
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

OUTDIR = sys.argv[1]
NAVY, BLUE, CYAN, RED = '08203F', '0A4FA0', '00AEEF', 'E2231A'
INK, MUTED, LINE, INPUT = '0B1B2E', '51647A', 'D5DEE8', 'EAF6FD'
F = 'Calibri'
font = lambda sz=10, b=False, c=INK, i=False: Font(name=F, size=sz, bold=b, color=c, italic=i)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = lambda c=LINE: Side(style='thin', color=c)

H = 13                      # ligne des en-têtes
R1, R2 = H + 1, H + 500     # lignes de saisie (formules déjà en place : ne jamais insérer de lignes, écrire dans la première ligne vide)
COLS = [  # (lettre, titre, largeur, type)
    ('A', 'Réf', 11, 'in'), ('B', 'Date\ndemande', 10.5, 'date'), ('C', 'Date\nenvoi', 10.5, 'date'), ('D', 'Délai\n(j)', 6, 'calc'),
    ('E', 'Client', 26, 'in'), ('F', 'Prospect /\nclient', 11, 'list'), ('G', 'Contact', 20, 'in'), ('H', 'Puits / rig', 18, 'in'), ('I', 'Job', 17, 'list'), ('J', 'Prestation', 11, 'list'), ('K', 'Description', 40, 'in'),
    ('L', 'Offre\npar', 8, 'in'), ('M', 'Incoterm', 11, 'list'), ('N', 'Statut', 13, 'list'), ('O', 'Date\nrelance', 10.5, 'date'), ('P', 'À relancer', 11, 'calc'),
    ('Q', 'Commentaires', 36, 'in'),
    ('R', 'Devis', 8, 'link'), ('S', 'Tool list', 9, 'link'),
    ('T', 'Nb\nBL', 5, 'in'), ('U', 'BL 1', 6, 'link'), ('V', 'BL 2', 6, 'link'), ('W', 'BL 3', 6, 'link'), ('X', 'BL 4', 6, 'link'),
    ('Y', 'BL 5', 6, 'link'), ('Z', 'BL 6', 6, 'link'), ('AA', 'BL 7', 6, 'link'), ('AB', 'BL 8', 6, 'link'),
    ('AC', 'Service\nticket', 9, 'link'), ('AD', 'Dossier', 9, 'link')]
PATH = 'Z:\\MUSTANG SERVICES\\12-Clients\\2027\\'   # dossier des affaires sur le serveur : Ctrl + H pour le remplacer dans les formules
LINKS = {'R': '_Devis.xlsx', 'S': '_ToolList.xlsx', 'AC': '_ServiceTicket.xlsx'}
BLS = {'U': 1, 'V': 2, 'W': 3, 'X': 4, 'Y': 5, 'Z': 6, 'AA': 7, 'AB': 8}   # BL n : Réf_BL.xlsx (n = 1), Réf_BL2.xlsx… Réf_BL8.xlsx, visible si « Nb BL » ≥ n
STATUTS = ['A faire', 'En validation', 'Envoyé', 'En attente', 'Gagné', 'Terminé', 'Perdu', 'Annulé']
JOBS = ['Fishing', 'Directional Drilling', 'Opérateurs', 'Serrage / desserrage', 'Fishing + DD', 'Machines', 'Autres']
PREST = ['Rental', 'Sales']
INCO = ['EXW Lons', 'FCA Anvers', 'FCA Norway', 'FOB Canada']

wb = Workbook(); ws = wb.active; ws.title = 'Chrono 2027'; ws.sheet_properties.tabColor = BLUE; ws.sheet_view.showGridLines = False
for l, t, w, k in COLS: ws.column_dimensions[l].width = w
LAST = len(COLS)
for r in (1, 2, 3):
    for c in range(1, LAST + 1): ws.cell(r, c).fill = fill(NAVY if r < 3 else CYAN)
ws.row_dimensions[1].height = 32; ws.row_dimensions[2].height = 20; ws.row_dimensions[3].height = 5
ws['A1'] = 'FRANCE CHRONO 2027 · ENEDRIL'; ws['A1'].font = font(18, True, 'FFFFFF'); ws['A1'].alignment = Alignment(vertical='center')
ws['A2'] = 'Un devis = une ligne. Écrire dans la première ligne vide du tableau (ne pas insérer de lignes : toutes les formules sont déjà en place).'
ws['A2'].font = font(10, c='9FD8F5'); ws['A2'].alignment = Alignment(vertical='center')

# --- indicateurs
def kpi(col, row, label, formula, fmt='0'):
    ws[f'{col}{row}'] = label; ws[f'{col}{row}'].font = font(8, True, MUTED); ws[f'{col}{row}'].alignment = Alignment(indent=1, vertical='bottom', wrap_text=True)
    ws[f'{col}{row+1}'] = formula; ws[f'{col}{row+1}'].font = font(15, True, BLUE); ws[f'{col}{row+1}'].number_format = fmt; ws[f'{col}{row+1}'].alignment = Alignment(horizontal='left', indent=1, vertical='center')
    ws[f'{col}{row+1}'].fill = fill(INPUT)
    ws[f'{col}{row}'].fill = fill(INPUT)
ST = f'$N${R1}:$N${R2}'
cn = lambda s: f'COUNTIF({ST},"{s}")'
ws.row_dimensions[4].height = 30; ws.row_dimensions[5].height = 26
for c1, c2, lab, f, fmt in (('A', 'A', 'Demandes', f'=COUNTA($A${R1}:$A${R2})', '0'), ('B', 'B', 'Devis faits', f'=COUNT($C${R1}:$C${R2})', '0'),
                            ('C', 'C', 'À faire / en validation', '=' + cn('A faire') + '+' + cn('En validation'), '0'),
                            ('E', 'E', 'En attente', '=' + cn('Envoyé') + '+' + cn('En attente'), '0'), ('F', 'F', 'Gagnés', '=' + cn('Gagné') + '+' + cn('Terminé'), '0'),
                            ('G', 'G', 'Perdus / annulés', '=' + cn('Perdu') + '+' + cn('Annulé'), '0'),
                            ('H', 'H', 'Taux de transformation (gagnés ÷ gagnés + perdus)', f'=IF(({cn("Gagné")}+{cn("Terminé")}+{cn("Perdu")})=0,"",({cn("Gagné")}+{cn("Terminé")})/({cn("Gagné")}+{cn("Terminé")}+{cn("Perdu")}))', '0%'),
                            ('I', 'I', "Durée moy. d'envoi (j)", f'=IFERROR(AVERAGE($D${R1}:$D${R2}),"")', '0.0'),
                            ('J', 'J', 'À relancer', f'=COUNTIF($P${R1}:$P${R2},"⚠*")', '0'),
                            ('K', 'K', '% Rental', f'=IF(COUNTA($J${R1}:$J${R2})=0,"",COUNTIF($J${R1}:$J${R2},"Rental")/COUNTA($J${R1}:$J${R2}))', '0%'),
                            ('L', 'L', '% Sales', f'=IF(COUNTA($J${R1}:$J${R2})=0,"",COUNTIF($J${R1}:$J${R2},"Sales")/COUNTA($J${R1}:$J${R2}))', '0%')):
    kpi(c1, 4, lab, f, fmt)
# --- répartition par type de job
JB = f'$I${R1}:$I${R2}'
WON = lambda: f'(COUNTIFS({ST},"Gagné")+COUNTIFS({ST},"Terminé"))'
ws['A7'] = 'RÉPARTITION PAR TYPE DE JOB'; ws['A7'].font = font(8.5, True, MUTED)
for r, lab in ((8, 'Job'), (9, '% des devis'), (10, '% des gagnés'), (11, 'Taux de transfo.')):
    ws.merge_cells(f'A{r}:D{r}'); ws[f'A{r}'] = lab; ws[f'A{r}'].font = font(8.5, True, NAVY)
    for cc in 'ABCD': ws[f'{cc}{r}'].fill = fill('DCE8F5')
    ws[f'A{r}'].alignment = Alignment(wrap_text=True, vertical='center', indent=1)
for i, job in enumerate(JOBS):
    col = L(5 + i)
    ws[f'{col}8'] = job; ws[f'{col}8'].font = font(8, True, 'FFFFFF'); ws[f'{col}8'].fill = fill(BLUE); ws[f'{col}8'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws[f'{col}9'] = f'=IF(COUNTA({JB})=0,"",COUNTIF({JB},{col}$8)/COUNTA({JB}))'
    ws[f'{col}10'] = f'=IF({WON()}=0,"",(COUNTIFS({JB},{col}$8,{ST},"Gagné")+COUNTIFS({JB},{col}$8,{ST},"Terminé"))/{WON()})'
    ws[f'{col}11'] = f'=IF(COUNTIFS({JB},{col}$8,{ST},"Gagné")+COUNTIFS({JB},{col}$8,{ST},"Terminé")+COUNTIFS({JB},{col}$8,{ST},"Perdu")=0,"",(COUNTIFS({JB},{col}$8,{ST},"Gagné")+COUNTIFS({JB},{col}$8,{ST},"Terminé"))/(COUNTIFS({JB},{col}$8,{ST},"Gagné")+COUNTIFS({JB},{col}$8,{ST},"Terminé")+COUNTIFS({JB},{col}$8,{ST},"Perdu")))'
    for r in (9, 10, 11):
        ws[f'{col}{r}'].number_format = '0%'; ws[f'{col}{r}'].font = font(10, True, BLUE); ws[f'{col}{r}'].fill = fill(INPUT); ws[f'{col}{r}'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[7].height = 16; ws.row_dimensions[8].height = 36
for r in (9, 10, 11): ws.row_dimensions[r].height = 18
ws.row_dimensions[12].height = 8
ws.row_dimensions[6].height = 8
# --- en-têtes
for l, t, w, k in COLS:
    c = ws[f'{l}{H}']; c.value = t; c.font = font(8.5, True, 'FFFFFF'); c.fill = fill(BLUE if k in ('in', 'date', 'list', 'eur') else ('2E8B57' if k == 'link' else '5B7390'))
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[H].height = 34
# --- lignes
for r in range(R1, R2 + 1):
    ws.row_dimensions[r].height = 18
    for l, t, w, k in COLS:
        c = ws[f'{l}{r}']; c.border = Border(bottom=thin('E6ECF3'), left=thin('F0F3F7'), right=thin('F0F3F7'))
        c.font = font(10, c=BLUE if k in ('in', 'date', 'list', 'eur') else INK); c.alignment = Alignment(vertical='center', horizontal='center' if k in ('date', 'calc', 'list', 'link') else None, indent=1 if l in 'EGHKQ' else 0)
        if k in ('date',): c.number_format = 'dd/mm/yy'
        if k == 'eur': c.number_format = '#,##0'; c.alignment = Alignment(horizontal='right', vertical='center', indent=1)
    ws[f'D{r}'] = f'=IF(AND(ISNUMBER($B{r}),ISNUMBER($C{r})),$C{r}-$B{r},"")'
    ws[f'P{r}'] = f'=IF(AND(OR($N{r}="Envoyé",$N{r}="En attente"),ISNUMBER($C{r}),NOT(ISNUMBER($O{r}))),IF(TODAY()-$C{r}>=7,"⚠ relancer",""),"")'
    for l, suf in LINKS.items():
        ws[f'{l}{r}'] = f'=IF($A{r}="","",HYPERLINK("{PATH}"&$A{r}&"\\"&$A{r}&"{suf}","Ouvrir"))'
        ws[f'{l}{r}'].font = Font(name=F, size=9.5, color='0563C1', underline='single')
    for l, n in BLS.items():
        suf = '_BL' + ('' if n == 1 else str(n)) + '.xlsx'
        ws[f'{l}{r}'] = f'=IF(OR($A{r}="",N($T{r})<{n}),"",HYPERLINK("{PATH}"&$A{r}&"\\"&$A{r}&"{suf}","BL {n}"))'
        ws[f'{l}{r}'].font = Font(name=F, size=9.5, color='0563C1', underline='single')
    ws[f'AD{r}'] = f'=IF($A{r}="","",HYPERLINK("{PATH}"&$A{r}&"\\","Dossier"))'; ws[f'AD{r}'].font = Font(name=F, size=9.5, color='0563C1', underline='single')
# listes déroulantes
def dv(col, items, strict=False):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True, showErrorMessage=strict); ws.add_data_validation(d); d.add(f'{col}{R1}:{col}{R2}')
dv('F', ['Prospect', 'Client']); dv('I', JOBS); dv('J', PREST); dv('M', INCO); dv('N', STATUTS, True)
# couleurs de statut
rng = f'N{R1}:N{R2}'
for s, bg, fg in (('A faire', 'E3E6EA', '4A5668'), ('Envoyé', 'B9D8F5', '0A4FA0'), ('En validation', 'E3D5F5', '4B2A7A'), ('En attente', 'FFE9A8', '7A5B00'), ('Gagné', '9EDDB0', '0B5A2A'), ('Terminé', 'CFE9D6', '0B5A2A'), ('Perdu', 'F6B3AE', '8A1710'), ('Annulé', 'D5DBE3', '4A5668')):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$N{R1}="{s}"'], fill=fill(bg), font=Font(name=F, bold=True, color=fg)))
ws.conditional_formatting.add(f'P{R1}:P{R2}', FormulaRule(formula=[f'LEFT($P{R1},1)="⚠"'], font=Font(name=F, bold=True, color=RED)))
# exemples fictifs
EX = [('27-001-FR', datetime.date(2027, 1, 5), datetime.date(2027, 1, 7), 'Client exemple A', 'Client', 'Contact A', 'Rig CLU108', 'Fishing', 'Rental', 'Exemple · Stab 5-7/8 + XO', 'PH', 'EXW Lons', 'Gagné'),
      ('27-002-CM', datetime.date(2027, 1, 8), datetime.date(2027, 1, 8), 'Client exemple B', 'Prospect', 'Contact B', 'Rig PAU-3', 'Autres', 'Sales', 'Exemple · Whipstock 10-3/4"', 'FJ/PH', 'EXW Lons', 'Envoyé'),
      ('27-003-DE', datetime.date(2027, 1, 12), None, 'Client exemple C', 'Client', 'Contact C', 'Rig LAC-12', 'Fishing', 'Rental', 'Exemple · Casing scraper', 'PH', 'EXW Lons', 'A faire')]
for i, row in enumerate(EX):
    r = R1 + i
    for l, v in zip('ABCEFGHIJKLMN', row):
        if v is not None: ws[f'{l}{r}'] = v
ws[f'T{R1}'] = 2; ws[f'T{R1+1}'] = 1   # exemple : 2 BL sur la 1re affaire, 1 sur la 2e
ws['Q' + str(R1)] = "LIGNES D'EXEMPLE (fictives) : à effacer"
ws['Q' + str(R1)].font = font(9, True, RED)
ws.freeze_panes = f'B{H+1}'; ws.auto_filter.ref = f'A{H}:{L(LAST)}{R2}'
ws.print_title_rows = f'{H}:{H}'; ws.print_area = f'A1:{L(LAST)}{R1 + 99}'
ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = .3; ws.page_margins.top = .4; ws.page_margins.bottom = .4
wb.calculation.fullCalcOnLoad = True
out = os.path.join(OUTDIR, 'France_Chrono_2027.xlsx'); wb.save(out); print('ok', out)
