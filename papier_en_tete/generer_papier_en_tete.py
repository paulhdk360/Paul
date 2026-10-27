# Papier en-tête Enedril A4 (corps vide) : python3 -I generer_papier_en_tete.py -> Papier_en_tete_Enedril_A4.docx
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
FONT='Ubuntu Mono';NAVY='08203F';BLUE='0A4FA0';GREY='5A6B7F'
LOGO='../certifications/logo_officiel_couleur_avec_baseline.png'
def setfont(run,size,color=None,bold=False):
    run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold
    rf=run._r.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):rf.set(qn(a),FONT)
    if color:run.font.color.rgb=RGBColor.from_string(color)
def border(p,side,color,sz=6,space=4):
    pPr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');e=OxmlElement('w:'+side)
    for k,v in (('val','single'),('sz',str(sz)),('space',str(space)),('color',color)):e.set(qn('w:'+k),v)
    b.append(e);pPr.append(b)
def fixed_table(container,widths_cm):
    t=container.add_table(1,len(widths_cm),Cm(sum(widths_cm)));t.autofit=False
    tblPr=t._tbl.tblPr
    lay=OxmlElement('w:tblLayout');lay.set(qn('w:type'),'fixed');tblPr.append(lay)
    tw=tblPr.find(qn('w:tblW'))
    if tw is None:tw=OxmlElement('w:tblW');tblPr.append(tw)
    tw.set(qn('w:w'),str(int(sum(widths_cm)*567)));tw.set(qn('w:type'),'dxa')
    grid=t._tbl.tblGrid
    for gc,w in zip(grid.findall(qn('w:gridCol')),widths_cm):gc.set(qn('w:w'),str(int(w*567)))
    for c,w in zip(t.rows[0].cells,widths_cm):c.width=Cm(w);c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    return t
d=Document();sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=sec.right_margin=Cm(2.2);sec.top_margin=Cm(4.2);sec.bottom_margin=Cm(3.6)
sec.header_distance=Cm(1.2);sec.footer_distance=Cm(1.0)
st=d.styles['Normal'];st.font.name=FONT;st.font.size=Pt(11.5);st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
W=16.6
# ---- EN-TÊTE : logo à gauche, coordonnées à droite, filet dessous
hd=sec.header;hd.paragraphs[0].paragraph_format.space_after=Pt(0)
t=fixed_table(hd,[7.6,W-7.6]);a,b=t.rows[0].cells
p=a.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.add_run().add_picture(LOGO,width=Cm(6.4))
lines=[('ENEDRIL',True,NAVY),('3 avenue des Lacs',False,GREY),('64140 Lons (France)',False,GREY),('+33 (0)5 59 04 50 45',False,GREY),('info@enedril.fr',False,GREY),('www.enedril.com',True,BLUE)]
for i,(tx,bd,col) in enumerate(lines):
    p=b.paragraphs[0] if i==0 else b.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(11.5)
    r=p.add_run(tx);setfont(r,9,col,bd)
p=hd.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(0);border(p,'top','1CB2FE',12,1)
r=p.add_run('');setfont(r,2)
# ---- PIED DE PAGE : filet, mentions légales à gauche, badges à droite
ft=sec.footer;fp=ft.paragraphs[0];fp.paragraph_format.space_after=Pt(0);border(fp,'top','D5E1EE',6,4)
r=fp.add_run('');setfont(r,2)
t=fixed_table(ft,[11.6,W-11.6]);a,b=t.rows[0].cells
for i,(tx,bd) in enumerate((('ENEDRIL · SAS au capital de 50 000 €',True),('SIRET 794 642 462 00020 · RCS Pau · Code APE 6420Z',False),('TVA intracommunautaire FR11794642462',False))):
    p=a.paragraphs[0] if i==0 else a.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(10.5)
    r=p.add_run(tx);setfont(r,8,NAVY if bd else GREY,bd)
p=b.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.space_after=Pt(0)
for f in ('../certifications/RSE_badge_entreprise_engagee.png','../certifications/LRQA_ISO9001_gris.png'):
    p.add_run().add_picture(f,height=Cm(1.45));p.add_run('  ')
# corps de la lettre : vide
d.save('Papier_en_tete_Enedril_A4.docx')
