# Flyer commercial Enedril A5 recto-verso : python3 -I generer_flyer.py  -> Flyer_Enedril_A5_FR/EN.pptx
import sys,qrcode
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
NAVY=(8,32,63);BLUE=(10,79,160);MID=(25,78,154);CYAN=(28,178,254);PALE=(234,244,252);INK=(11,27,46);GREY=(90,107,127);WHITE=(255,255,255);LINE=(213,225,238)
F='Ubuntu Mono';W,H=420,595
LOGO='../public/logo-enedril-white.png';TOOL='../outil_SPX_bleu.png'
T={
'fr':dict(h1=['FORAGE DIRIGÉ','& FISHING'],h2='Solutions pour la géothermie, le pétrole et gaz et le stockage souterrain de gaz.',
  stats=[('98 500 m','forés'),('75','puits, jusqu\'à 3 000 m TVD'),('142','interventions fishing, workover et P&A'),('2 500','outils en stock')],
  iso='ISO 9001:2015',since='Depuis 2013 à Lons, près de Pau',
  met='NOS MÉTIERS',svc=[('FORAGE DIRIGÉ & MWD','Moteurs de fond 3-1/8" à 9-5/8", RSS, MWD jusqu\'à 200°C, planification de puits en interne.'),('FISHING','Tiges, trains de tiges, outils coincés. Réactivité 24/48 h.'),('WORKOVER','Section milling, wash over, réparation de casing.'),('BOUCHAGE & ABANDON (P&A)','Abandon de puits en toute sécurité.')],
  why='POURQUOI ENEDRIL',whyl=['EXPÉRIENCE','EXCELLENCE','SUR MESURE','FIABILITÉ'],
  contact='Parlons de votre projet',addr='3 avenue des Lacs, 64140 Lons (France)',scan='Scannez pour nous découvrir'),
'en':dict(h1=['DIRECTIONAL DRILLING','& FISHING'],h2='Solutions for geothermal, oil & gas and underground gas storage projects.',
  stats=[('98,500 m','drilled'),('75','wells, up to 3,000 m TVD'),('142','fishing, workover and P&A jobs'),('2,500','tools in stock')],
  iso='ISO 9001:2015',since='Since 2013 in Lons, near Pau (France)',
  met='WHAT WE DO',svc=[('DIRECTIONAL DRILLING & MWD','PDM 3-1/8" to 9-5/8", RSS, MWD up to 200°C, in-house well planning.'),('FISHING','Rods, strings, stuck tools. Response 24/48 h.'),('WORKOVER','Section milling, wash over, casing repair.'),('PLUG & ABANDONMENT','Safe abandonment of wells.')],
  why='WHY ENEDRIL',whyl=['EXPERIENCE','EXCELLENCE','CUSTOM SOLUTIONS','RELIABILITY'],
  contact='Let\'s talk about your project',addr='3 avenue des Lacs, 64140 Lons (France)',scan='Scan to discover us')}
def build(lang):
    t=T[lang];P=Presentation();P.slide_width=Pt(W);P.slide_height=Pt(H)
    def poly(s,pts,fill):
        fb=s.shapes.build_freeform(Pt(pts[0][0]),Pt(pts[0][1]));fb.add_line_segments([(Pt(x),Pt(y)) for x,y in pts[1:]],close=True)
        sh=fb.convert_to_shape();sh.shadow.inherit=False;sh.fill.solid();sh.fill.fore_color.rgb=C(*fill);sh.line.fill.background();return sh
    rect=lambda s,x,y,w,h,c:poly(s,[(x,y),(x+w,y),(x+w,y+h),(x,y+h)],c)
    def txt(s,x,y,w,h,text,size=10,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,line=None):
        tb=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=tb.text_frame;tf.word_wrap=True;tf.vertical_anchor=anchor
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,x_ in enumerate(text if isinstance(text,list) else [text]):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
            if line:p.line_spacing=Pt(line)
            r=p.add_run();r.text=x_;r.font.size=Pt(size);r.font.bold=bold;r.font.name=F;r.font.color.rgb=C(*color)
    # ---- recto
    s=P.slides.add_slide(P.slide_layouts[6]);rect(s,0,0,W,H,NAVY)
    poly(s,[(0,300),(W,200),(W,H),(0,H)],BLUE);poly(s,[(0,395),(W,330),(W,H),(0,H)],MID)
    tl=s.shapes.add_picture(TOOL,Pt(-95),Pt(105),width=Pt(560));tl.rotation=-24
    s.shapes.add_picture(LOGO,Pt(28),Pt(28),width=Pt(130))
    txt(s,28,252,370,20,t['since'].upper(),8.5,CYAN,True)
    for i,l in enumerate(t['h1']):txt(s,28,270+i*36,380,42,l,(30 if max(len(z) for z in t['h1'])<15 else 25),WHITE,True)
    txt(s,28,350,330,50,t['h2'],10.5,(200,222,245),line=14.5)
    for i,(v,l) in enumerate(t['stats']):
        x=28+(i%2)*190;y=405+(i//2)*68
        rect(s,x,y,3,48,CYAN);txt(s,x+12,y-2,170,28,v,22,WHITE,True);txt(s,x+12,y+26,170,26,l,8.5,(190,215,240),line=11)
    rect(s,28,H-38,110,20,WHITE);txt(s,28,H-38,110,20,t['iso'],9,NAVY,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    txt(s,150,H-38,240,20,'www.enedril.com',9.5,CYAN,True,PP_ALIGN.RIGHT,MSO_ANCHOR.MIDDLE)
    # ---- verso
    s=P.slides.add_slide(P.slide_layouts[6]);rect(s,0,0,W,H,WHITE)
    poly(s,[(260,0),(W,0),(W,95),(215,95)],PALE);poly(s,[(0,385),(W,345),(W,H),(0,H)],NAVY)
    txt(s,28,28,300,24,t['met'],20,NAVY,True);rect(s,28,58,44,4,CYAN)
    for i,(k,d) in enumerate(t['svc']):
        y=80+i*68;rect(s,28,y,3,52,CYAN);txt(s,42,y-1,350,16,k,11.5,BLUE,True);txt(s,42,y+17,350,40,d,9.5,INK,line=13)
    txt(s,28,398,300,16,t['why'],10,CYAN,True)
    for i,w in enumerate(t['whyl']):
        x=28+i*95;rect(s,x,420,85,3,CYAN);
    # (why : mots-clés sous le titre, sur le fond marine)
    ov=s.shapes[-1]
    for i,w in enumerate(t['whyl']):txt(s,28+i*95,428,92,14,w,8,WHITE,True)
    txt(s,28,470,230,18,t['contact'],14,WHITE,True)
    for i,l in enumerate(['+33 (0)5 59 04 50 45','info@enedril.fr','www.enedril.com',t['addr']]):txt(s,28,498+i*19,260,16,l,10,(200,222,245) if i==3 else WHITE,i<3)
    q=qrcode.QRCode(border=1,box_size=10);q.add_data('https://www.enedril.com');q.make();im=q.make_image(fill_color='#08203F',back_color='white').convert('RGB');im.save('qr_enedril.png')
    rect(s,296,470,100,100,WHITE);s.shapes.add_picture('qr_enedril.png',Pt(301),Pt(475),width=Pt(90))
    txt(s,290,574,112,12,t['scan'],6.5,(190,215,240),align=PP_ALIGN.CENTER)
    P.save(f'Flyer_Enedril_A5_{lang.upper()}.pptx')
for l in ('fr','en'):build(l)
