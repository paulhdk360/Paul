# Plan marketing & communication 2026 : python3 -I generer_plan.py  -> Plan_marketing_communication_2026_v3.xlsx
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.formatting.rule import DataBarRule,ColorScaleRule,FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,LineChart,Reference,Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as L
NAVY='08203F';BLUE='0A4FA0';CYAN='1CB2FE';SKY='01ADEE';PALE='EAF4FC';PALE2='D3E5F5';INK='0B1B2E';GREY='5A6B7F';WHITE='FFFFFF';LINE='D5E1EE';INPUT='0000FF'
fill=lambda c:PatternFill('solid',fgColor=c)
thin=Side(style='thin',color=LINE);BOT=Border(bottom=thin);BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def f(sz=10,b=False,c=INK,i=False,name='Calibri'):return Font(name=name,size=sz,bold=b,color=c,italic=i)
AL=lambda h='left',v='center',w=False:Alignment(horizontal=h,vertical=v,wrap_text=w)
MOIS=['Janvier','Février','Mars','Avril','Mai','Juin','Juillet','Août','Septembre','Octobre','Novembre','Décembre']
wb=Workbook();D=wb.active;D.title='Tableau de bord'
CA=wb.create_sheet('Calendrier');PO=wb.create_sheet('Posts LinkedIn');AB=wb.create_sheet('Abonnés LinkedIn');SA=wb.create_sheet('Salons')
for ws,col in ((D,NAVY),(CA,BLUE),(PO,SKY),(AB,CYAN),(SA,GREY)):ws.sheet_properties.tabColor=col;ws.sheet_view.showGridLines=False
def banner(ws,title,sub,last):
    for r in (1,2,3):
        for c in range(1,last+1):ws.cell(r,c).fill=fill(NAVY)
    ws['B1'].value=title;ws['B1'].font=f(22,True,WHITE,name='Segoe UI');ws['B2'].value=sub;ws['B2'].font=f(10,False,CYAN)
    ws.row_dimensions[1].height=34;ws.row_dimensions[2].height=18;ws.row_dimensions[3].height=8
    for c in range(1,last+1):ws.cell(4,c).fill=fill(CYAN)
    ws.row_dimensions[4].height=3
def head(ws,r,c1,labels,h=30):
    for i,t in enumerate(labels):
        c=ws.cell(r,c1+i,t);c.font=f(10,True,WHITE);c.fill=fill(BLUE);c.alignment=AL('center','center',True)
    ws.row_dimensions[r].height=h
PL="'Posts LinkedIn'!";AL_="'Abonnés LinkedIn'!";CAL="Calendrier!"
# ================= POSTS =================
ws=PO;banner(ws,'Posts LinkedIn 2026','Saisissez uniquement les cellules en bleu : taux, coûts et mois se calculent tout seuls.',17)
cols=['N°','Post','Date','Mois','Type','Statut','Vues vidéo','Impressions','Engagements',"Taux d'engagement",'Clics','Taux de clics','Réactions','Commentaires','Partages','Budget (€)','Coût / 1000 impr.']
widths=[6,34,13,7,11,11,11,13,13,13,9,11,11,13,10,11,13]
for i,w in enumerate(widths):ws.column_dimensions[L(i+1)].width=w
ws['B5']='Total publiés';ws['B5'].font=f(10,True,NAVY)
R0,R1=7,66
for c in range(1,18):ws.cell(5,c).fill=fill(PALE);ws.cell(5,c).font=f(10,True,NAVY)
for col in 'GHIKMNOP':ws[col+'5']=f'=SUM({col}{R0}:{col}{R1})';ws[col+'5'].number_format='#,##0'
ws['J5']='=IFERROR(I5/H5,"")';ws['L5']='=IFERROR(K5/H5,"")';ws['J5'].number_format=ws['L5'].number_format='0.0%'
ws['Q5']='=IFERROR(P5/SUMIFS(H%d:H%d,P%d:P%d,">0")*1000,"")'%(R0,R1,R0,R1);ws['Q5'].number_format='0.00 "€"'
ws['F5']=f'=COUNTIF(F{R0}:F{R1},"Publié")&" posts"';ws['A5']=''
head(ws,6,1,cols,34)
posts=[('Vœux 2026',(2026,1,6),'Post',None,1525,87,34,52,1,0,0),('Arrivée Koen',(2026,2,10),'Post',None,4120,348,258,81,9,0,0),('Annonce présence GeoTHERM',(2026,2,20),'Post',None,945,114,60,46,7,1,0),('Enedril au GeoTHERM',(2026,2,26),'Post',None,5339,1211,1043,158,6,4,0),('DTS350',(2026,3,12),'Vidéo',3196,6252,642,515,120,2,None,0),('Performance Kem One',(2026,3,24),'Post',None,7985,2196,2049,132,13,2,0),('Junk Mill GMX 10',(2026,4,2),'Post',None,2665,955,876,75,3,1,0),('HIRING : Assistante comptable',(2026,4,13),'Post',None,1338,101,75,18,0,8,0),('Post 4000 abonnés',(2026,4,23),'Post',None,22024,168,97,60,2,6,200),('Cartes Jobs paris Geothermie',(2026,5,21),'Post',None,2447,268,166,93,4,5,0),('Annonce Geodays',(2026,6,11),'Post',None,883,42,10,31,0,1,0),('Retour sur Geodays',(2026,6,26),'Post',None,1160,154,61,90,2,1,0),('Soletanche',(2026,10,1),'Post',None,1907,333,259,62,0,3,0),('Teambuilding',(2026,10,6),'Post',None,1838,1120,1071,52,0,0,0),('Salon GEOFLUID',(2026,10,8),'Post',None,763,45,17,26,0,2,0),('Octobre Rose',(2026,10,13),'Post',None,2166,269,208,64,0,0,0),('Prospection foreurs',(2026,10,15),'Post',None,47945,508,435,32,1,1,140),('Réception IFP',(2026,10,20),'Post',None,1122,402,372,29,1,0,0)]
plan=['Case study part 1','Case study part 2','Anti-collision KEM One','Nouvelle Break Out']
for i in range(R1-R0+1):
    r=R0+i;ws.cell(r,1,i+1).font=f(9,False,GREY);ws.cell(r,1).alignment=AL('center')
    if i<len(posts):
        t,d,ty,v,imp,eng,clk,rea,com,par,bud=posts[i]
        vals={2:t,3:dt.date(*d),5:ty,6:'Publié',7:v,8:imp,9:eng,11:clk,13:rea,14:com,15:par,16:bud}
    elif i-len(posts)<len(plan):vals={2:plan[i-len(posts)],5:'Post',6:'Planifié'}
    else:vals={}
    for c,v in vals.items():ws.cell(r,c,v)
    ws.cell(r,4,f'=IF(C{r}="","",MONTH(C{r}))');ws.cell(r,10,f'=IF(AND(H{r}>0,I{r}<>""),I{r}/H{r},"")');ws.cell(r,12,f'=IF(AND(H{r}>0,K{r}<>""),K{r}/H{r},"")')
    ws.cell(r,17,f'=IF(AND(N(P{r})>0,N(H{r})>0),P{r}/H{r}*1000,"")')
    for c in range(1,18):
        x=ws.cell(r,c);x.border=BOT;x.alignment=AL('left' if c==2 else 'center')
        if c in (2,3,5,6,7,8,9,11,13,14,15,16):x.font=f(10,False,INPUT)
        elif c!=1:x.font=f(10)
        if r%2==0:x.fill=fill('F7FAFD')
    ws.cell(r,3).number_format='dd/mm/yyyy'
    for c in (7,8,9,11,13,14,15):ws.cell(r,c).number_format='#,##0'
    ws.cell(r,10).number_format=ws.cell(r,12).number_format='0.0%';ws.cell(r,16).number_format='0 "€"';ws.cell(r,17).number_format='0.00 "€"'
dv=DataValidation(type='list',formula1='"Post,Vidéo,Carrousel,Article,Sondage,Événement"',allow_blank=True);dv.add(f'E{R0}:E{R1}');ws.add_data_validation(dv)
dv2=DataValidation(type='list',formula1='"Publié,Planifié,Idée,Annulé"',allow_blank=True);dv2.add(f'F{R0}:F{R1}');ws.add_data_validation(dv2)
ws.conditional_formatting.add(f'H{R0}:H{R1}',DataBarRule(start_type='num',start_value=0,end_type='max',color=CYAN))
ws.conditional_formatting.add(f'J{R0}:J{R1}',ColorScaleRule(start_type='min',start_color='FFFFFF',end_type='max',end_color='5CB8F5'))
ws.conditional_formatting.add(f'L{R0}:L{R1}',ColorScaleRule(start_type='min',start_color='FFFFFF',end_type='max',end_color='5CB8F5'))
ws.conditional_formatting.add(f'F{R0}:F{R1}',FormulaRule(formula=[f'F{R0}="Planifié"'],font=Font(color='B26A00',bold=True)))
ws.conditional_formatting.add(f'F{R0}:F{R1}',FormulaRule(formula=[f'F{R0}="Publié"'],font=Font(color='1B7F3B',bold=True)))
ws.freeze_panes='C7'
# ================= ABONNES =================
ws=AB;banner(ws,'Abonnés LinkedIn','Nouveaux abonnés par mois : saisissez le mois en cours, le cumul et les KPI se mettent à jour.',17)
ws.column_dimensions['A'].width=3;ws.column_dimensions['B'].width=10
for c in range(3,15):ws.column_dimensions[L(c)].width=10
ws.column_dimensions['O'].width=10;ws.column_dimensions['P'].width=3;ws.column_dimensions['Q'].width=42;ws.column_dimensions['R'].width=14
ws['B6']='NOUVEAUX ABONNÉS PAR MOIS';ws['B6'].font=f(11,True,BLUE)
head(ws,7,2,['Année']+[m[:4]+'.' if len(m)>4 else m for m in MOIS]+['Total'],24)
data={2024:[37,11,50,47,186,169,144,153,125,206,129,79],2025:[132,102,144,108,191,162,88,113,93,108,65,40],2026:[47,91,124,142,125,68,70,23,68,123,None,None]}
for i,y in enumerate((2024,2025,2026)):
    r=8+i;ws.cell(r,2,y).font=f(10,True,NAVY)
    for m,v in enumerate(data[y]):
        c=ws.cell(r,3+m,v);c.font=f(10,False,INPUT);c.alignment=AL('center');c.border=BOT
    t=ws.cell(r,15,f'=SUM(C{r}:N{r})');t.font=f(10,True,NAVY);t.fill=fill(PALE);t.alignment=AL('center')
ws['B12']='ABONNÉS CUMULÉS';ws['B12'].font=f(11,True,BLUE)
head(ws,13,2,['Année']+[m[:4]+'.' if len(m)>4 else m for m in MOIS],24)
anch={2024:'$R$8',2025:'$R$9',2026:'$R$10'}
for i,y in enumerate((2024,2025,2026)):
    r=14+i;ws.cell(r,2,y).font=f(10,True,NAVY);src=8+i
    for m in range(12):
        cl=L(3+m);c=ws.cell(r,3+m,f'=IF({cl}{src}="",NA(),{anch[y]}+SUM($C{src}:{cl}{src}))');c.alignment=AL('center');c.font=f(10);c.number_format='#,##0';c.border=BOT
ws.conditional_formatting.add('C14:N16',FormulaRule(formula=['ISNA(C14)'],font=Font(color=WHITE)))

ws['Q6']='REPÈRES ET OBJECTIF';ws['Q6'].font=f(11,True,BLUE)
kp=[('Abonnés au 01/01/2024',1067,'#,##0'),('Abonnés au 31/12/2024 (relevé)',2397,'#,##0'),('Abonnés au 31/12/2025 (relevé)',3743,'#,##0'),('Abonnés actuels','=R10+SUM(C10:N10)','#,##0'),('Croissance 2026 (à date)','=R11-R10','+#,##0'),('Croissance 2025','=R10-R9','+#,##0'),('Nouveaux abonnés / mois en 2026 (moyenne)','=IFERROR(AVERAGE(C10:N10),0)','0'),('Objectif 2026',5000,'#,##0'),('Restant à atteindre','=MAX(R15-R11,0)','#,##0'),("% de l'objectif atteint",'=R11/R15','0%'),('Projection fin 2026 (rythme actuel)','=R11+R14*(12-COUNT(C10:N10))','#,##0'),('Même période 2025 (nouveaux abonnés)','=SUMPRODUCT((C10:N10<>"")*C9:N9)','#,##0'),('Évolution vs même période 2025','=IFERROR(R12/R19-1,"")','+0%;-0%')]
for i,(lab,v,fm) in enumerate(kp):
    r=8+i;a=ws.cell(r,17,lab);a.font=f(10);a.border=BOT;b=ws.cell(r,18,v);b.number_format=fm;b.alignment=AL('right');b.border=BOT
    isin=not str(v).startswith('=');b.font=f(11,True,INPUT if isin else NAVY)
ws['R17'].fill=fill(PALE);ws.conditional_formatting.add('R17',DataBarRule(start_type='num',start_value=0,end_type='num',end_value=1,color=CYAN))
ws['Q22']='Note : le relevé au 31/12/2024 (2 397) est inférieur de 6 au calcul (1 067 + 1 336 = 2 403), l’écart correspond probablement à des désabonnements.';ws['Q22'].font=f(9,False,GREY,True);ws['Q22'].alignment=AL('left','top',True);ws.merge_cells('Q22:R25')
c1=LineChart();c1.title='Abonnés cumulés';c1.height=8;c1.width=17
for i,(y,col) in enumerate(((2024,'9DB6D4'),(2025,BLUE),(2026,CYAN))):
    se=Series(Reference(ws,min_col=3,max_col=14,min_row=14+i),title=str(y));se.graphicalProperties.line.solidFill=col;se.graphicalProperties.line.width=28000 if y==2026 else 19000;se.smooth=False;c1.series.append(se)
c1.set_categories(Reference(ws,min_col=3,max_col=14,min_row=13));c1.legend.position='b';ws.add_chart(c1,'B19')
c2=BarChart();c2.type='col';c2.title='Nouveaux abonnés par mois';c2.height=8;c2.width=17
for i,(y,col) in enumerate(((2024,'9DB6D4'),(2025,BLUE),(2026,CYAN))):
    se=Series(Reference(ws,min_col=3,max_col=14,min_row=8+i),title=str(y));se.graphicalProperties.solidFill=col;c2.series.append(se)
c2.set_categories(Reference(ws,min_col=3,max_col=14,min_row=7));c2.legend.position='b';ws.add_chart(c2,'B36')
ws.freeze_panes='A5'
# ================= CALENDRIER =================
import openpyxl
src=openpyxl.load_workbook('source_original_v2.xlsx')['Calendrier']
def clean(v):
    if not v:return ''
    return '\n'.join('• '+x.replace('►','').strip() for x in str(v).split('\n') if x.strip())
ws=CA;banner(ws,'Calendrier éditorial 2026','LinkedIn, blog et événements mois par mois. Le mois en cours est surligné.',14)
ws.column_dimensions['A'].width=2;ws.column_dimensions['B'].width=22
for c in range(3,15):ws.column_dimensions[L(c)].width=27
head(ws,6,2,['']+MOIS,26)
ws['B5']='';
labels=['Posts prévus','Posts publiés','Écart','Impressions','Engagements','Thèmes LinkedIn','Blog','Salons et événements']
heights=[22,22,22,22,22,190,60,95]
for i,(lab,h) in enumerate(zip(labels,heights)):
    r=7+i;c=ws.cell(r,2,lab);c.font=f(10,True,NAVY);c.fill=fill(PALE);c.alignment=AL('left','center',True);ws.row_dimensions[r].height=h;c.border=BOT
prev=[src.cell(7,2+m).value for m in range(12)]
for m in range(12):
    col=3+m;cl=L(col)
    ws.cell(7,col,prev[m]).font=f(11,False,INPUT)
    ws.cell(8,col,f'=COUNTIFS({PL}$D$7:$D$66,COLUMN()-2,{PL}$F$7:$F$66,"Publié")')
    ws.cell(9,col,f'=IF(COLUMN()-2>MONTH(TODAY()),"",{cl}8-{cl}7)')
    ws.cell(10,col,f'=SUMIFS({PL}$H$7:$H$66,{PL}$D$7:$D$66,COLUMN()-2)');ws.cell(11,col,f'=SUMIFS({PL}$I$7:$I$66,{PL}$D$7:$D$66,COLUMN()-2)')
    for r in range(7,12):
        x=ws.cell(r,col);x.alignment=AL('center');x.border=BOT
        if r>7:x.font=f(11,r==8)
        if r in (10,11):x.number_format='#,##0'
    ws.cell(12,col,clean(src.cell(8,2+m).value));ws.cell(13,col,'À planifier');ws.cell(14,col,clean(src.cell(22,2+m).value))
    for r in (12,13,14):
        x=ws.cell(r,col);x.alignment=AL('left','top',True);x.font=f(10,False,GREY if r==13 else INK,r==13);x.border=Border(left=thin,bottom=thin)
ws.conditional_formatting.add('C6:N6',FormulaRule(formula=['COLUMN()-2=MONTH(TODAY())'],fill=fill(CYAN),font=Font(color=NAVY,bold=True)))
ws.conditional_formatting.add('C7:N14',FormulaRule(formula=['COLUMN()-2=MONTH(TODAY())'],fill=fill('E3F3FE')))
ws.conditional_formatting.add('C9:N9',FormulaRule(formula=['AND(C9<>"",C9<0)'],font=Font(color='C0392B',bold=True)))
ws.conditional_formatting.add('C9:N9',FormulaRule(formula=['AND(C9<>"",C9>=0)'],font=Font(color='1B7F3B',bold=True)))
ws.freeze_panes='C7'
# ================= SALONS =================
ws=SA;banner(ws,'Salons et événements','Suivi des salons confirmés, du lieu et du nombre de posts LinkedIn liés (calcul automatique via le mot-clé).',9)
for c,w in zip('ABCDEFGHI',[2,34,16,22,13,26,12,46,18]):ws.column_dimensions[c].width=w
head(ws,6,2,['Événement','Type','Période','Statut','Lieu','Posts liés','Remarque','Mot-clé'],26)
rows=[('Journées de la Géothermie','Salon','Février 2026','À valider','','',''),('GeoTHERM 2026','Salon','Février 2026','Réalisé','Offenburg (Allemagne)','Annonce et retour publiés.','GeoTHERM'),('Think GeoEnergy','Salon','Mars 2026','À valider','','',''),('GeoRin','Salon','Avril 2026','À valider','','',''),('WPC Energy Congress','Salon','Avril 2026','À valider','','',''),('World Hydrogen Summit & Exhibition','Salon','Mai 2026','À valider','','',''),('Geodays','Salon','Juin 2026','Réalisé','','Annonce et retour publiés.','Geodays'),('GEOFLUID','Salon','9 octobre 2026','Réalisé','Piacenza (Italie)','Représentant : Luigi Liberatore.','GEOFLUID'),('Journée mondiale de l’énergie','Temps fort','16 octobre 2026','Réalisé','','','énergie'),('Octobre rose','Temps fort','Octobre 2026','Réalisé','','','Octobre Rose'),('ADIPEC 2026','Salon','2-5 novembre 2026','Confirmé','Abu Dhabi (ADNEC)','Post d’annonce à publier avant le 2 novembre.','ADIPEC'),('Movember','Temps fort','Novembre 2026','Confirmé','','','Movember'),('GeoTHERM 2027','Salon','25-26 février 2027','Confirmé','Offenburg (Allemagne)','Stand 274 (18 m²), co-exposant Fixndril.','GeoTHERM 2027')]
for i,(n,ty,pe,st,li,re,kw) in enumerate(rows):
    r=7+i;vals=[n,ty,pe,st,li,None,re,kw]
    for j,v in enumerate(vals):
        c=ws.cell(r,2+j,v);c.border=BOT;c.alignment=AL('left' if j in (0,4,6) else 'center','center',True);c.font=f(10,j==0,INPUT if j in (2,3,4,6,7) else INK)
        if r%2==0:c.fill=fill('F7FAFD')
    ws.cell(r,7,f'=IF(I{r}="","",COUNTIF({PL}$B$7:$B$66,"*"&I{r}&"*"))').font=f(11,True,NAVY)
    ws.row_dimensions[r].height=24
dv=DataValidation(type='list',formula1='"Confirmé,Réalisé,Annulé,À valider"',allow_blank=True);dv.add('E7:E40');ws.add_data_validation(dv)
dv=DataValidation(type='list',formula1='"Salon,Temps fort,Conférence"',allow_blank=True);dv.add('C7:C40');ws.add_data_validation(dv)
ws.conditional_formatting.add('E7:E40',FormulaRule(formula=['E7="Réalisé"'],fill=fill('D6F0DE'),font=Font(color='1B7F3B',bold=True)))
ws.conditional_formatting.add('E7:E40',FormulaRule(formula=['E7="Confirmé"'],fill=fill('DCEBFA'),font=Font(color=BLUE,bold=True)))
ws.conditional_formatting.add('E7:E40',FormulaRule(formula=['E7="Annulé"'],fill=fill('F8D7D3'),font=Font(color='C0392B',bold=True)))
ws.freeze_panes='C7'
# ================= TABLEAU DE BORD =================
ws=D;banner(ws,'Plan marketing et communication 2026','Enedril · tableau de bord LinkedIn, calendrier et salons. Tout se met à jour depuis les autres onglets.',14)
ws.column_dimensions['A'].width=2
for c in range(2,14):ws.column_dimensions[L(c)].width=13.5
ws.column_dimensions['N'].width=2
def card(col,row,label,formula,fmt,sub,color=BLUE,big=22):
    for r in range(row,row+4):
        for c in range(col,col+3):ws.cell(r,c).fill=fill(WHITE if False else PALE)
    for r in range(row,row+4):ws.cell(r,col).border=Border(left=Side(style='thick',color=color))
    ws.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+2);ws.merge_cells(start_row=row+1,start_column=col,end_row=row+2,end_column=col+2);ws.merge_cells(start_row=row+3,start_column=col,end_row=row+3,end_column=col+2)
    a=ws.cell(row,col,label.upper());a.font=f(8.5,True,GREY);a.alignment=AL('left','center');a.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    b=ws.cell(row+1,col,formula);b.font=f(big,True,color,name='Segoe UI');b.number_format=fmt;b.alignment=Alignment(horizontal='left',vertical='center',indent=1,wrap_text=True)
    c=ws.cell(row+3,col,sub);c.font=f(9,False,GREY);c.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[row].height=18;ws.row_dimensions[row+3].height=18
P_=lambda col:f"{PL}${col}$7:${col}$66"
card(2,6,'Abonnés LinkedIn',f"={AL_}R11",'#,##0',f'=TEXT({AL_}R17,"0%")&" de l\'objectif ("&TEXT({AL_}R15,"#,##0")&")"',BLUE)
card(5,6,'Nouveaux abonnés 2026',f"={AL_}R12",'+#,##0',f'=IFERROR(TEXT({AL_}R20,"+0%;-0%")&" vs même période 2025","")',CYAN)
card(8,6,'Posts publiés 2026',f'=COUNTIF({P_("F")},"Publié")','0',f'=COUNTIF({P_("F")},"Planifié")&" planifiés · "&TEXT(COUNTIF({P_("F")},"Publié")/MAX(MONTH(TODAY()),1),"0.0")&" posts / mois"',NAVY)
card(11,6,'Impressions',f'={PL}H5','#,##0',f'=TEXT({PL}H5/MAX(COUNTIF({P_("F")},"Publié"),1),"#,##0")&" par post en moyenne"',BLUE)
card(2,11,'Engagements',f'={PL}I5','#,##0',f'=TEXT({PL}M5,"#,##0")&" réactions · "&{PL}N5&" commentaires"',CYAN)
card(5,11,"Taux d'engagement",f'={PL}J5','0.0%','=IFERROR("Taux de clics : "&TEXT('+PL+'L5,"0.0%"),"")',NAVY)
card(8,11,'Budget publicitaire',f'={PL}P5','#,##0 "€"',f'=IFERROR("Coût : "&TEXT({PL}Q5,"0.00")&" € / 1000 impressions","")',BLUE)
card(11,11,'Partages',f'={PL}O5','#,##0',f'=TEXT({PL}K5,"#,##0")&" clics au total"',CYAN)
card(2,16,"Meilleur taux d'engagement",f'=IFERROR(INDEX({P_("B")},MATCH(MAX({P_("J")}),{P_("J")},0)),"")','@',f'=IFERROR(TEXT(MAX({P_("J")}),"0%")&" d\'engagement","")',NAVY,13)
card(5,16,'Post le plus vu',f'=IFERROR(INDEX({P_("B")},MATCH(MAX({P_("H")}),{P_("H")},0)),"")','@',f'=TEXT(MAX({P_("H")}),"#,##0")&" impressions"',BLUE,13)
card(8,16,'Salons réalisés',f'=COUNTIFS(Salons!$C$7:$C$40,"Salon",Salons!$E$7:$E$40,"Réalisé")','0',f'=COUNTIFS(Salons!$C$7:$C$40,"Salon",Salons!$E$7:$E$40,"Confirmé")&" salons à venir"',CYAN)
card(11,16,'À compléter',f'=COUNTIFS({P_("F")},"Publié",{P_("H")},"")+COUNTIFS({P_("F")},"Publié",{P_("I")},"")','0','posts publiés sans statistiques','C0392B')
def put(ch,anchor):ws.add_chart(ch,anchor)
def style(ch,t,w=15.2,h=7.4):ch.title=t;ch.width=w;ch.height=h;ch.legend=None
c=BarChart();c.type='col';style(c,'Impressions par post');se=Series(Reference(PO,min_col=8,min_row=7,max_row=26),title='Impressions');se.graphicalProperties.solidFill=CYAN;c.series.append(se);c.set_categories(Reference(PO,min_col=2,min_row=7,max_row=26));c.gapWidth=60;put(c,'B22')
c=BarChart();c.type='col';style(c,"Taux d'engagement par post");se=Series(Reference(PO,min_col=10,min_row=7,max_row=26),title='Taux');se.graphicalProperties.solidFill=BLUE;c.series.append(se);c.set_categories(Reference(PO,min_col=2,min_row=7,max_row=26));c.gapWidth=60;c.y_axis.number_format='0%';put(c,'H22')
c=BarChart();c.type='col';style(c,'Posts prévus et publiés par mois',h=7.4);c.legend=None
for r,col in ((7,'9DB6D4'),(8,CYAN)):
    se=Series(Reference(CA,min_col=3,max_col=14,min_row=r),title=CA.cell(r,2).value);se.graphicalProperties.solidFill=col;c.series.append(se)
c.set_categories(Reference(CA,min_col=3,max_col=14,min_row=6));c.legend=__import__('openpyxl').chart.legend.Legend();c.legend.position='b';put(c,'B37')
c=BarChart();c.type='col';style(c,'Impressions par mois');se=Series(Reference(CA,min_col=3,max_col=14,min_row=10),title='Impressions');se.graphicalProperties.solidFill=BLUE;c.series.append(se);c.set_categories(Reference(CA,min_col=3,max_col=14,min_row=6));put(c,'H37')
c=LineChart();style(c,'Abonnés cumulés');
for i,(y,col) in enumerate(((2024,'9DB6D4'),(2025,BLUE),(2026,CYAN))):
    se=Series(Reference(AB,min_col=3,max_col=14,min_row=14+i),title=str(y));se.graphicalProperties.line.solidFill=col;se.graphicalProperties.line.width=28000 if y==2026 else 19000;se.smooth=False;c.series.append(se)
c.set_categories(Reference(AB,min_col=3,max_col=14,min_row=13));c.legend=__import__('openpyxl').chart.legend.Legend();c.legend.position='b';put(c,'B52')
c=BarChart();c.type='col';style(c,'Nouveaux abonnés par mois')
for i,(y,col) in enumerate(((2024,'9DB6D4'),(2025,BLUE),(2026,CYAN))):
    se=Series(Reference(AB,min_col=3,max_col=14,min_row=8+i),title=str(y));se.graphicalProperties.solidFill=col;c.series.append(se)
c.set_categories(Reference(AB,min_col=3,max_col=14,min_row=7));c.legend=__import__('openpyxl').chart.legend.Legend();c.legend.position='b';put(c,'H52')
# historique annuel
r0=68;ws.cell(r0,2,'ÉVOLUTION ANNUELLE').font=f(11,True,BLUE)
head(ws,r0+1,2,['Indicateur','','2024','2025','2026 (à date)','Évol. 24→25','Évol. 25→26'],26)
ws.merge_cells(start_row=r0+1,start_column=2,end_row=r0+1,end_column=3)
hist=[('Posts publiés',29,35,f'=COUNTIF({P_("F")},"Publié")','0'),('Impressions',187339,648545,f'={PL}H5','#,##0'),('Engagements',27672,27877,f'={PL}I5','#,##0'),("Taux d'engagement",'=D{e}/D{i}','=E{e}/E{i}','=F{e}/F{i}','0.0%'),('Clics',17886,20958,f'={PL}K5','#,##0'),('Taux de clics','=D{c}/D{i}','=E{c}/E{i}','=F{c}/F{i}','0.0%'),('Réactions',1882,2675,f'={PL}M5','#,##0'),('Commentaires',96,167,f'={PL}N5','#,##0'),('Partages',50,96,f'={PL}O5','#,##0'),('Budget (€)',960,2500,f'={PL}P5','#,##0')]
rowof={n:r0+2+i for i,(n,*_) in enumerate(hist)}
for i,(n,a,b,cv,fm) in enumerate(hist):
    r=r0+2+i;ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3);ws.cell(r,2,n).font=f(10,True,NAVY)
    for j,v in enumerate((a,b,cv)):
        if isinstance(v,str):v=v.format(e=rowof['Engagements'],i=rowof['Impressions'],c=rowof['Clics'])
        x=ws.cell(r,4+j,v);x.number_format=fm;x.alignment=AL('center');x.font=f(10,False,INPUT if j<2 and not str(v).startswith('=') else INK)
    if fm=='0.0%':
        ws.cell(r,7,f'=IFERROR((E{r}-D{r})*100,"")').number_format='+0.0" pts";-0.0" pts"';ws.cell(r,8,f'=IFERROR((F{r}-E{r})*100,"")').number_format='+0.0" pts";-0.0" pts"'
    else:
        ws.cell(r,7,f'=IFERROR(E{r}/D{r}-1,"")').number_format='+0%;-0%';ws.cell(r,8,f'=IFERROR(F{r}/E{r}-1,"")').number_format='+0%;-0%'
    for cc in range(2,9):ws.cell(r,cc).border=BOT
    for cc in (7,8):ws.cell(r,cc).alignment=AL('center');ws.cell(r,cc).font=f(10,True)
ws.conditional_formatting.add(f'G{r0+2}:H{r0+11}',FormulaRule(formula=[f'AND(ISNUMBER(G{r0+2}),G{r0+2}>0)'],font=Font(color='1B7F3B',bold=True)))
ws.conditional_formatting.add(f'G{r0+2}:H{r0+11}',FormulaRule(formula=[f'AND(ISNUMBER(G{r0+2}),G{r0+2}<0)'],font=Font(color='C0392B',bold=True)))
ws.cell(r0+13,2,'2026 est comparé à des années complètes : les évolutions 25→26 augmenteront avec les derniers mois. Valeurs 2024 et 2025 saisies à la main (cellules bleues).').font=f(9,False,GREY,True)
for w_ in wb:
    w_.sheet_view.zoomScale=100;w_.page_setup.orientation='landscape';w_.page_setup.fitToWidth=1;w_.page_setup.fitToHeight=0;w_.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
wb.save('Plan_marketing_communication_2026_v3.xlsx')
