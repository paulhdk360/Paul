import sys; sys.path.insert(0,'.')
from tr_fr import fr
from pptx import Presentation
SRC='/root/.claude/uploads/25d99196-61ed-597d-8b44-fc98194d121f/5b74efb0-Enedril_-_Presentation_-_EN_-_paysage.pptx'
p=Presentation(SRC); miss=set()
def tf_do(tf):
    for pa in tf.paragraphs:
        for r in pa.runs:
            if not r.text.strip(): continue
            t=fr(r.text)
            if t is None:
                if not r.text.strip().isdigit(): miss.add(r.text)
            else:
                r.text=t
            rPr=r._r.get_or_add_rPr(); rPr.set('lang','fr-FR')
def walk(shapes):
    for sh in shapes:
        if sh.shape_type==6: walk(sh.shapes)
        if sh.has_text_frame: tf_do(sh.text_frame)
        if getattr(sh,'has_table',False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells: tf_do(c.text_frame)
for s in p.slides: walk(s.shapes)
print('missing',miss)
p.save('/home/user/Paul/Presentation_Enedril_FR_paysage.pptx')
